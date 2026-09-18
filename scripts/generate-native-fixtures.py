"""Render the validated synthetic fixture sources; never read user documents."""
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

root = Path(__file__).resolve().parents[1] / 'output/isolation/current_runs/20260907_business_pipeline/fixtures'
if root.resolve() != root:
    raise ValueError('Fixture directory must not be a symlink')
purchase = json.loads((root / 'procurement.txt').read_text())
audit = json.loads((root / 'audit.txt').read_text())
for source in (purchase, audit):
    assert source['testFlag'] == 'MOCK_TEST_DATA_DO_NOT_USE'
subtotal = sum(row['quantity'] * row['unitPriceMinor'] for row in purchase['items'])
assert subtotal == purchase['financials']['subtotalMinor']
assert purchase['financials']['totalMinor'] == subtotal + purchase['financials']['taxMinor']
assert audit['differenceMinor'] == sum(row['bookMinor'] - row['checkedMinor'] for row in audit['entries'])

def money(minor):
    return f"USD {minor // 100:,}.{minor % 100:02d}"

document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.8)
document.styles['Normal'].font.name = 'Arial'
document.styles['Normal'].font.size = Pt(10)
for style in document.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
document.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
document.core_properties.author = 'Synthetic Test Fixture'
document.core_properties.last_modified_by = 'Synthetic Test Fixture'
document.core_properties.created = document.core_properties.modified = datetime(2026, 9, 7, tzinfo=timezone.utc)
document.add_paragraph('Procurement Contract Review', 'Title')
document.add_paragraph('MOCK TEST DATA ONLY - NO CONTRACTUAL EFFECT')
document.add_paragraph('Management asks the selected sales reviewers to check the items, prices and delivery date. This review does not authorize signing or payment.')
document.add_paragraph(f"Document: {purchase['documentId']}   Version: 1   Date: 2026-09-07")
document.add_heading('Parties and purchase items', level=1)
document.add_paragraph(f"Buyer: {purchase['buyer']}\nSupplier: {purchase['seller']}")
table = document.add_table(rows=1, cols=4)
table.style = 'Light Shading Accent 1'
for cell, text in zip(table.rows[0].cells, ['Item', 'Quantity', 'Unit price', 'Line total']):
    cell.text = text
for row in purchase['items']:
    for cell, value in zip(table.add_row().cells, [row['name'], str(row['quantity']), money(row['unitPriceMinor']), money(row['quantity'] * row['unitPriceMinor'])]):
        cell.text = value
for label, key in [('Subtotal', 'subtotalMinor'), ('Synthetic tax at 5 percent', 'taxMinor'), ('Total', 'totalMinor')]:
    document.add_paragraph(label + ': ' + money(purchase['financials'][key]))
document.add_heading('Delivery and payment', level=1)
document.add_paragraph(f"Delivery date: {purchase['deliveryDate']}\nDelivery destination: {purchase['deliveryMethod']}\nPayment terms: {purchase['paymentTerms']}")
document.add_heading('Review limits', level=1)
for key in ['confidentiality', 'changeCondition', 'breachHandling']:
    document.add_paragraph(purchase[key])
raw = io.BytesIO()
document.save(raw)
normalized = io.BytesIO()
with zipfile.ZipFile(raw) as source, zipfile.ZipFile(normalized, 'w', zipfile.ZIP_DEFLATED) as target:
    for name in sorted(source.namelist()):
        info = zipfile.ZipInfo(name, date_time=(2026, 9, 7, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        target.writestr(info, source.read(name))

pdf = io.BytesIO()
styles = getSampleStyleSheet()
story = [Paragraph('Internal Audit Reconciliation', styles['Title']),
         Paragraph('MOCK TEST DATA ONLY - NO FINANCIAL AUTHORITY', styles['Heading2']),
         Paragraph('Accounting requests an internal audit review of synthetic expense discrepancies. The selected auditor may report findings but cannot authorize payment or asset seizure.', styles['BodyText']),
         Spacer(1, 14), Paragraph(f"Document: {audit['documentId']} | Period: {audit['auditPeriod']} | Version: 1", styles['BodyText']),
         Paragraph('Expense reconciliation', styles['Heading2'])]
rows = [['Voucher', 'Category', 'Book', 'Checked']]
for row in audit['entries']:
    rows.append([row['voucherId'], row['category'], money(row['bookMinor']), money(row['checkedMinor'])])
table = Table(rows, colWidths=[108, 125, 100, 100], repeatRows=1)
table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e6eeee')),
                          ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                          ('FONTSIZE', (0, 0), (-1, -1), 9),
                          ('TOPPADDING', (0, 0), (-1, -1), 9),
                          ('BOTTOMPADDING', (0, 0), (-1, -1), 9)]))
story.extend([table, Spacer(1, 12), Paragraph('Difference: ' + money(audit['differenceMinor']), styles['Heading2']),
              Paragraph('Missing document: ' + audit['missingDocuments'][0], styles['BodyText']),
              Paragraph('Action: ' + audit['pendingItems'][0], styles['BodyText']),
              Paragraph('Review deadline: ' + audit['dueAt'], styles['BodyText']),
              Paragraph('Authority limit', styles['Heading2']), Paragraph(audit['responsibility'], styles['BodyText'])])
SimpleDocTemplate(pdf, pagesize=A4, rightMargin=54, leftMargin=54, topMargin=48, bottomMargin=48,
                  title='Synthetic Internal Audit', author='Synthetic Test Fixture', invariant=1).build(story)
outputs = {'procurement.docx': normalized.getvalue(), 'audit.pdf': pdf.getvalue()}
for name, payload in outputs.items():
    destination = root / name
    if destination.is_symlink() or (destination.exists() and destination.read_bytes() != payload):
        raise ValueError('Existing native fixture differs; refusing overwrite')
for name, payload in outputs.items():
    destination = root / name
    if not destination.exists():
        with destination.open('xb') as stream:
            stream.write(payload)
        destination.chmod(0o600)
    print(f'{name}: {len(payload)} bytes; isolation only')
